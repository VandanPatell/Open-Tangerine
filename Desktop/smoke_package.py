import csv,json,os,subprocess,tempfile
from pathlib import Path
from zipfile import ZipFile
from PIL import Image
from openpyxl import Workbook,load_workbook
from pypdf import PdfReader
root=Path(__file__).resolve().parent
exe=root/'dist/OpenTangerine'/('OpenTangerine.exe' if os.name=='nt' else 'OpenTangerine')
subprocess.run([str(exe),'--smoke-test'],check=True,timeout=30)

def worker(path,target,folder):
 r=subprocess.run([str(exe),'--worker'],input=json.dumps({'paths':[str(path)],'target':target,'output':str(folder)}),text=True,capture_output=True,timeout=30)
 # Keep console-mode packaging: worker JSON needs real standard streams.
 if r.returncode!=0:raise RuntimeError(r.stderr or r.stdout)
 data=json.loads(r.stdout);assert data['ok'];return [Path(p) for p in data['outputs']]

with tempfile.TemporaryDirectory() as folder:
 folder=Path(folder)
 p=folder/'input.png';Image.new('RGB',(20,20),'orange').save(p)
 out=worker(p,'JPEG',folder)[0]
 with Image.open(out) as im:assert im.size==(20,20)
 spreadsheet=folder/'workbook.xlsx';book=Workbook()
 book.active.append(['name','value']);book.active.append(['café','00123'])
 book.create_sheet('Second').append(['second sheet']);book.save(spreadsheet);book.close()
 exports=worker(spreadsheet,'Excel → CSV',folder);assert len(exports)==2
 with exports[0].open(encoding='utf-8-sig',newline='') as stream:
  assert list(csv.reader(stream))==[['name','value'],['café','00123']]
 restored=worker(exports[0],'CSV → Excel',folder)[0]
 book=load_workbook(restored)
 try:assert book.active['B2'].value=='00123'
 finally:book.close()
 document=folder/'document.docx'
 with ZipFile(document,'w') as archive:
  archive.writestr('word/document.xml','<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Packaged café document</w:t></w:r></w:p></w:body></w:document>')
 text=worker(document,'DOCX → Text',folder)[0]
 assert text.read_text(encoding='utf-8')=='Packaged café document'
 pdf=PdfReader(worker(document,'DOCX → PDF (text)',folder)[0])
 assert 'Packaged café document' in pdf.pages[0].extract_text()
 fonts=pdf.pages[0]['/Resources']['/Font'].get_object().values()
 assert any('/FontFile2' in font.get_object().get('/FontDescriptor',{}).get_object()
            for font in fonts if '/FontDescriptor' in font.get_object()), 'Bundled PDF font missing'
print('PASS: packaged startup, image, multi-sheet CSV, XLSX, DOCX text/PDF and embedded font')
