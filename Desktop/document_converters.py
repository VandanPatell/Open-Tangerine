"""Bounded local spreadsheet and text-document conversions; no Office automation."""
import csv
import os
from pathlib import Path
import re
import subprocess
import tempfile
import time
from xml.sax.saxutils import escape
from zipfile import ZipFile

from engine import ConversionError, MAX_OUTPUT, publish

OFFICE_INPUTS = {'.doc', '.docx', '.odt', '.rtf', '.xls', '.xlsx', '.ods', '.ppt', '.pptx', '.odp'}
MAX_EXPANDED = 128 * 1024 * 1024
MAX_CELLS = 1_000_000
MAX_TEXT = 2 * 1024 * 1024


def check_package(path):
    # Inspect members without extracting them; reject ZIP bombs before XML parsing.
    with ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > 10000 or sum(m.file_size for m in members) > MAX_EXPANDED:
            raise ConversionError('Office document exceeds the 128 MiB expanded/10,000 part limit.')
        if any(m.flag_bits & 1 for m in members):
            raise ConversionError('Encrypted Office documents are not supported.')


def excel_csv(path, work):
    from openpyxl import load_workbook
    check_package(path)
    # No macro execution, external link loading, or formula recalculation.
    values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    formulas = None
    pieces = []
    try:
        formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        if not 1 <= len(values.worksheets) <= 100:
            raise ConversionError('Workbook must contain 1–100 worksheets.')
        cells = 0
        for index, sheet in enumerate(values.worksheets):
            if (sheet.max_row or 0) > 100000 or (sheet.max_column or 0) > 512:
                raise ConversionError('CSV export supports up to 100,000 rows and 512 columns per sheet.')
            output = work / f'sheet-{index}.csv'
            with output.open('w', newline='', encoding='utf-8-sig') as stream:
                writer = csv.writer(stream)
                for row, formula_row in zip(sheet.iter_rows(), formulas.worksheets[index].iter_rows()):
                    cells += len(row)
                    if cells > MAX_CELLS:
                        raise ConversionError('Workbook exceeds the one-million-cell export limit.')
                    exported = []
                    for cell, formula in zip(row, formula_row):
                        if formula.data_type == 'f' and cell.value is None:
                            raise ConversionError('A formula has no cached result. Recalculate and save in Excel before CSV export.')
                        value = cell.value
                        # Keep CSV values as data when subsequently opened in Excel.
                        if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
                            value = "'" + value
                        exported.append(value)
                    writer.writerow(exported)
                    if stream.tell() > MAX_OUTPUT:
                        raise ConversionError('CSV output exceeds 512 MiB.')
            if output.stat().st_size <= 3:
                raise ConversionError('Workbook contains an empty worksheet; no CSVs were published.')
            pieces.append((output, f'{path.stem}-{index+1}-{sheet.title}', 'csv'))
            if sum(p.stat().st_size for p, _, _ in pieces) > MAX_OUTPUT:
                raise ConversionError('Combined CSV output exceeds 512 MiB.')
        return pieces
    finally:
        values.close()
        if formulas is not None: formulas.close()


def csv_excel(path, work):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    book = Workbook(write_only=True)
    sheet = book.create_sheet('Data')
    cells = 0
    try:
        with path.open(encoding='utf-8-sig', newline='') as stream:
            for index, row in enumerate(csv.reader(stream)):
                cells += len(row)
                if index >= 100000 or len(row) > 512 or cells > MAX_CELLS:
                    raise ConversionError('CSV exceeds 100,000 rows, 512 columns or one million cells.')
                if any(len(value) > 32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value) for value in row):
                    raise ConversionError('CSV contains a cell Excel cannot represent without truncation.')
                literals = []
                for value in row:
                    cell = WriteOnlyCell(sheet, value=value)
                    cell.data_type = 's'  # Never turn user text into executable formulas.
                    literals.append(cell)
                sheet.append(literals)
        if not cells: raise ConversionError('CSV contains no cells.')
        output = work / 'table.xlsx'
        book.save(output)
        return [(output, path.stem+'-converted', 'xlsx')]
    finally:
        if not sheet.closed: sheet.close()
        book.close()


def docx_paragraphs(path):
    from defusedxml.ElementTree import fromstring
    check_package(path)
    with ZipFile(path) as archive:
        info = archive.getinfo('word/document.xml')
        if info.file_size > 16 * 1024 * 1024:
            raise ConversionError('Document XML exceeds 16 MiB.')
        root = fromstring(archive.read(info))
    ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    paragraphs = []
    length = 0
    for paragraph in root.iter(ns+'p'):
        text = ''.join((node.text or '') if node.tag == ns+'t' else '\t' if node.tag == ns+'tab'
                       else '\n' if node.tag in (ns+'br', ns+'cr') else '' for node in paragraph.iter())
        length += len(text)
        if length > MAX_TEXT or len(paragraphs) >= 10000:
            raise ConversionError('Document exceeds the two-million-character/10,000 paragraph limit.')
        paragraphs.append(text)
    if not any(p.strip() for p in paragraphs):
        raise ConversionError('Document has no extractable body text.')
    return paragraphs


def text_pdf(paragraphs, output):
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas
    font = 'Helvetica'
    # ReportLab's bundled Vera provides an embedded font on every supported OS.
    import reportlab
    font_path = Path(reportlab.__file__).parent / 'fonts/Vera.ttf'
    if font_path.is_file():
        pdfmetrics.registerFont(TTFont('TangerineText', str(font_path)))
        font = 'TangerineText'
        # Fail instead of silently replacing unsupported characters with boxes.
        supported = pdfmetrics.getFont(font).face.charToGlyph
        if any(ord(c) not in supported for line in paragraphs for c in line if not c.isspace()):
            raise ConversionError('Text uses characters unavailable in the PDF font. Use a document editor or LibreOffice for multilingual PDF export.')
    style = ParagraphStyle('Body', fontName=font, fontSize=11, leading=16, spaceAfter=8)
    story = []
    for line in paragraphs:
        # Escape ReportLab markup so documents cannot add links/images or load URLs.
        story.append(Paragraph(escape(line).replace('\n', '<br/>').replace('\t', '    '), style)
                     if line.strip() else Spacer(1, 12))

    class BoundedCanvas(Canvas):
        def showPage(self):
            if self.getPageNumber() > 200:
                raise ConversionError('Text PDF exceeds 200 pages.')
            super().showPage()

    SimpleDocTemplate(str(output), pagesize=(595.28, 841.89),
                      leftMargin=48, rightMargin=48, topMargin=48, bottomMargin=48).build(story, canvasmaker=BoundedCanvas)


def office_convert(path, work, executable, extension='pdf'):
    exe = Path(executable or '')
    if not executable or not exe.is_absolute() or not exe.is_file() or exe.is_symlink():
        raise ConversionError('Select a trusted LibreOffice executable for formatted Office PDF conversion.')
    profile = work / 'office-profile'; profile.mkdir()
    # A private profile avoids sending documents to an existing interactive session.
    (profile/'user').mkdir()
    (profile/'user/registrymodifications.xcu').write_text('''<?xml version="1.0"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item></oor:items>''', encoding='utf-8')
    out = work/'office-output'; out.mkdir()
    with tempfile.TemporaryFile() as log:
        proc = subprocess.Popen([str(exe), '-env:UserInstallation='+profile.as_uri(), '--headless',
            '--norestore', '--nodefault', '--convert-to', extension, '--outdir', str(out), str(path)],
            stdin=subprocess.DEVNULL, stdout=log, stderr=log,
            **({'creationflags': subprocess.CREATE_NO_WINDOW} if os.name == 'nt' else {'start_new_session': True}))
        started = time.monotonic()
        try:
            while proc.poll() is None:
                if time.monotonic()-started > 120 or log.tell() > 65536 or sum(p.stat().st_size for p in out.glob('*') if p.is_file()) > MAX_OUTPUT:
                    raise ConversionError('LibreOffice exceeded conversion limits.')
                time.sleep(.05)
            output = out/(path.stem+'.'+extension)
            if proc.returncode or not output.is_file():
                raise ConversionError('LibreOffice could not convert this document. Check its format or password protection.')
            return [(output, path.stem+'-converted', extension)]
        finally:
            if proc.poll() is None:
                from runner import stop
                stop(proc)


def convert_document(path, target, work, folder, office=''):
    suffix = path.suffix.lower()
    if target == 'Excel → CSV' and suffix == '.xlsx':
        pieces = excel_csv(path, work)
    elif target == 'Excel → CSV' and suffix in ('.xls', '.ods'):
        if suffix == '.ods': check_package(path)
        intermediate = office_convert(path, work, office, 'xlsx')[0][0]
        pieces = excel_csv(intermediate, work)
    elif target == 'CSV → Excel' and suffix == '.csv':
        pieces = csv_excel(path, work)
    elif target in ('DOCX → Text', 'DOCX → PDF (text)') and suffix == '.docx':
        paragraphs = docx_paragraphs(path)
        output = work/'document'
        if target == 'DOCX → Text':
            output.write_text('\n'.join(paragraphs), encoding='utf-8'); ext = 'txt'
        else:
            text_pdf(paragraphs, output); ext = 'pdf'
        pieces = [(output, path.stem+'-converted', ext)]
    elif target == 'Text → PDF' and suffix in ('.txt', '.md'):
        if path.stat().st_size > MAX_TEXT:
            raise ConversionError('Text PDF input must be UTF-8 and no larger than 2 MiB.')
        text = path.read_text(encoding='utf-8-sig')
        if not text.strip(): raise ConversionError('Text document is empty.')
        output = work/'document.pdf'; text_pdf(text.splitlines(), output)
        pieces = [(output, path.stem+'-converted', 'pdf')]
    elif target == 'Office → PDF' and suffix in OFFICE_INPUTS:
        if suffix in ('.xlsx', '.docx', '.pptx', '.odt', '.ods', '.odp'): check_package(path)
        pieces = office_convert(path, work, office)
    else:
        raise ConversionError('This operation requires a matching document or spreadsheet input.')
    # Parse all sheets before publishing any output; existing outputs are preserved.
    if any(not p.is_file() or not 0 < p.stat().st_size <= MAX_OUTPUT for p, _, _ in pieces):
        raise ConversionError('Empty or oversized document output.')
    return [publish(p, folder, name, ext) for p, name, ext in pieces]
