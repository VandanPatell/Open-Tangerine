import csv
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch, Mock
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engine import convert, ConversionError
from runner import run_job
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def job(self, path, target, **options):
        return dict(paths=[str(path)], target=target, output=str(self.root), **options)

    def docx(self, text='Hello &amp; welcome'):
        path = self.root/'sample.docx'
        with ZipFile(path, 'w') as archive:
            archive.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>'+text+'</w:t></w:r></w:p><w:p><w:r><w:t>Second paragraph</w:t></w:r></w:p></w:body></w:document>')
        return path

    def test_excel_all_sheets_to_csv_unicode_and_no_overwrites(self):
        path = self.root/'book.xlsx'
        book = Workbook(); book.active.title = 'Sales'
        book.active.append(['Name', 'Amount']); book.active.append(['café, "tea"', 12.5])
        book.create_sheet('Second').append(['=literal', '00123'])
        book['Second']['A1'].data_type = 's'
        book.save(path); book.close()
        original = path.read_bytes()
        result = run_job(self.job(path, 'Excel → CSV'))
        self.assertEqual(len(result), 2)
        with Path(result[0]).open(encoding='utf-8-sig', newline='') as stream:
            self.assertEqual(list(csv.reader(stream)), [['Name', 'Amount'], ['café, "tea"', '12.5']])
        self.assertIn("'=literal", Path(result[1]).read_text(encoding='utf-8-sig'))
        again = convert(self.job(path, 'Excel → CSV'))
        self.assertFalse(set(result) & set(again))
        self.assertEqual(path.read_bytes(), original)

    def test_uncached_formula_fails_without_publishing_any_sheets(self):
        path = self.root/'formula.xlsx'
        book = Workbook(); book.active.append(['first sheet'])
        book.create_sheet('Formula').append(['=1+2']); book.save(path); book.close()
        with self.assertRaisesRegex(ConversionError, 'cached result'):
            convert(self.job(path, 'Excel → CSV'))
        self.assertEqual(list(self.root.glob('*.csv')), [])

    def test_csv_to_excel_preserves_literal_formulas_and_leading_zeroes(self):
        path = self.root/'table.csv'
        path.write_text('name,code,formula\n"café, tea",00123,=1+2\n', encoding='utf-8-sig')
        output = run_job(self.job(path, 'CSV → Excel'))[0]
        book = load_workbook(output)
        try:
            self.assertEqual(book.active['A2'].value, 'café, tea')
            self.assertEqual(book.active['B2'].value, '00123')
            self.assertEqual(book.active['C2'].value, '=1+2')
            self.assertEqual(book.active['C2'].data_type, 's')
        finally: book.close()

    def test_csv_limits_and_invalid_encoding(self):
        path = self.root/'bad.csv'; path.write_text('x'*32768)
        with self.assertRaisesRegex(ConversionError, 'truncation'):
            convert(self.job(path, 'CSV → Excel'))
        path.write_bytes(b'\xff\xff')
        with self.assertRaises(UnicodeError): convert(self.job(path, 'CSV → Excel'))
        self.assertFalse(list(self.root.glob('*.xlsx')))

    def test_docx_text_and_pdf_worker(self):
        path = self.docx()
        original = path.read_bytes()
        text = run_job(self.job(path, 'DOCX → Text'))[0]
        self.assertEqual(Path(text).read_text(encoding='utf-8'), 'Hello & welcome\nSecond paragraph')
        pdf = run_job(self.job(path, 'DOCX → PDF (text)'))[0]
        extracted = ''.join(page.extract_text() for page in PdfReader(pdf).pages)
        self.assertIn('Hello & welcome', extracted)
        self.assertIn('Second paragraph', extracted)
        self.assertEqual(path.read_bytes(), original)

    def test_document_markup_never_loads_remote_images(self):
        path = self.root/'text.txt'
        path.write_text('<img src="https://example.com/secret"/>\nLiteral text', encoding='utf-8')
        pdf = convert(self.job(path, 'Text → PDF'))[0]
        text = ''.join(p.extract_text() for p in PdfReader(pdf).pages)
        self.assertIn('<img src=', text)

    def test_docx_entities_and_expanded_limits_rejected(self):
        path = self.root/'malicious.docx'
        with ZipFile(path, 'w') as archive:
            archive.writestr('word/document.xml', '<!DOCTYPE x [<!ENTITY x SYSTEM "file:///etc/passwd">]><x>&x;</x>')
        with self.assertRaises(Exception): convert(self.job(path, 'DOCX → Text'))
        with patch('document_converters.MAX_EXPANDED', 5):
            with self.assertRaisesRegex(ConversionError, 'expanded'):
                convert(self.job(self.docx(), 'DOCX → Text'))
        self.assertFalse(list(self.root.glob('*.txt')))

    def test_unsupported_pdf_glyph_fails_without_loss(self):
        path = self.docx('Hello &#128512;')
        with self.assertRaisesRegex(ConversionError, 'characters unavailable'):
            convert(self.job(path, 'DOCX → PDF (text)'))
        self.assertFalse(list(self.root.glob('*.pdf')))

    def test_pdf_page_limit(self):
        from document_converters import text_pdf
        with self.assertRaisesRegex(ConversionError, '200 pages'):
            text_pdf(['line']*12000, self.root/'oversized.pdf')

    def test_office_requires_local_executable(self):
        with self.assertRaisesRegex(ConversionError, 'LibreOffice'):
            convert(self.job(self.docx(), 'Office → PDF'))

    def test_office_private_profile_and_argument_arrays(self):
        path = self.docx()
        def launch(arguments, **kwargs):
            output = Path(arguments[arguments.index('--outdir')+1]) / (path.stem+'.pdf')
            from document_converters import text_pdf
            text_pdf(['Office fixture'], output)
            self.assertIn('--headless', arguments)
            self.assertEqual(arguments[-1], str(path))
            from urllib.parse import urlsplit
            from urllib.request import url2pathname
            profile = Path(url2pathname(urlsplit(arguments[1].split('=',1)[1]).path))
            self.assertIn('DisableMacrosExecution', (profile/'user/registrymodifications.xcu').read_text())
            return Mock(poll=Mock(return_value=0), returncode=0)
        with patch('document_converters.subprocess.Popen', side_effect=launch):
            self.assertTrue(convert(self.job(path, 'Office → PDF', office=str(Path(sys.executable).resolve()))))

    def test_office_timeout_stops_its_process(self):
        process = Mock(poll=Mock(return_value=None))
        with patch('document_converters.subprocess.Popen', return_value=process), \
             patch('document_converters.time.monotonic', side_effect=[0, 121]), \
             patch('runner.stop') as stop:
            with self.assertRaisesRegex(ConversionError, 'limits'):
                convert(self.job(self.docx(), 'Office → PDF', office=str(Path(sys.executable).resolve())))
            stop.assert_called_once_with(process)


if __name__ == '__main__': unittest.main()
